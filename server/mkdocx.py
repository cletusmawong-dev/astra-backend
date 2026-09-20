import json, sys
from docx import Document
from docx.shared import Pt, RGBColor

spec = json.load(open(sys.argv[1], encoding="utf-8"))
doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
t = doc.add_heading(spec["title"], level=0)
for sec in spec["sections"]:
    doc.add_heading(sec["h"], level=1)
    if sec.get("p"):
        doc.add_paragraph(sec["p"])
    for b in sec.get("b", []):
        doc.add_paragraph(b, style="List Bullet")
doc.save(spec["out"])
print("ok")
