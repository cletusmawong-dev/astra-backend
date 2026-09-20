import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

spec = json.load(open(sys.argv[1], encoding="utf-8"))
wb = Workbook()
ws = wb.active
ws.title = spec["title"][:28] or "Sheet"
hdr = Font(bold=True, color="FFFFFF")
fill = PatternFill("solid", fgColor="4F8CFF")
for c, h in enumerate(spec["headers"], 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.font = hdr
    cell.fill = fill
for r, row in enumerate(spec["rows"], 2):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
for c in range(1, len(spec["headers"]) + 1):
    ws.column_dimensions[chr(64 + c)].width = 18
wb.save(spec["out"])
print("ok")
